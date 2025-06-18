import json
import logging
import os
from datetime import datetime

from django.contrib.auth.decorators import login_required
from django.shortcuts import render, get_object_or_404, redirect
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
from django.db import transaction
from django.utils.translation import gettext as _
from django.utils.translation import gettext_lazy as _l

from .forms import DocumentUploadForm, MedicalEventForm, CustomUserCreationForm
from .models import Document, MedicalEvent, LabIndicator, BloodTestMeasurement, NarrativeSectionResult, DocumentTag, \
    Practitioner, PatientProfile, MedicalCategory, MedicalSpecialty
from ocrapi.vision_handler import perform_ocr_space
from ocrapi.gpt_client import call_gpt_for_document
from django.urls import reverse
from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Font
import hashlib
from ocrapi import anonymizer

logger = logging.getLogger(__name__)

ALLOWED_IMAGE_EXTENSIONS = ['.jpg', '.jpeg', '.png', '.gif', '.tiff', '.bmp', '.webp']
ALLOWED_PDF_EXTENSIONS = ['.pdf']


# --- Public-facing pages ---
def landing_page(request):
    return render(request, 'basetemplates/landingpage.html')


def register_page(request):
    if request.method == 'POST':
        form = CustomUserCreationForm(request.POST)
        if form.is_valid():
            user = form.save()
            PatientProfile.objects.create(user=user)
            return redirect('medj:login')
    else:
        form = CustomUserCreationForm()
    return render(request, 'registration/register.html', {'form': form})


# --- Authenticated pages ---
@login_required
def dashboard_page(request):
    return render(request, 'main/dashboard.html')


@login_required
def casefiles_page(request):
    return render(request, 'main/casefiles.html')


@login_required
def personalcard_page(request):
    return render(request, 'main/personalcard.html')


@login_required
def upload_page(request):
    messages = {
        'select_file_type': _("Моля, изберете тип на файла."),
        'select_file': _("Моля, изберете файл за качване."),
        'file_too_large': _("Размерът на файла е твърде голям. Максимален размер: 10 MB."),
        'invalid_file_type': _("Невалиден тип файл. Поддържани формати: JPG, JPEG, PNG, GIF, TIFF, BMP, WEBP, PDF."),
        'uploading_file': _("Качва се файл..."),
        'ocr_processing': _("Обработка с OCR, моля изчакайте..."),
        'analysis_error': _("Грешка при анализ:"),
        'critical_analysis_error': _("Критична грешка при анализ:"),
        'ocr_success': _("OCR обработката приключи успешно."),
        'approve_analyze_ai': _("Одобри и анализирай с AI"),
        'analyze_loading': _("Извършва се AI анализ, моля изчакайте..."),
        'upload_success': _("Файлът е качен успешно!"),
        'upload_error': _("Грешка при качване:"),
        'critical_upload_error': _("Критична грешка при качване:"),
    }

    event_types = MedicalEvent.EventType.choices
    categories = MedicalCategory.objects.all().order_by('name')
    specialties = MedicalSpecialty.objects.all().order_by('name')
    doctors = Practitioner.objects.all().order_by('name')

    context = {
        'MESSAGES': json.dumps(messages),
        'event_types': event_types,
        'categories': categories,
        'specialties': specialties,
        'doctors': doctors,
    }

    return render(request, 'main/upload.html', context)


@login_required
def upload_history_page(request):
    documents = Document.objects.filter(patient=request.user.patientprofile).order_by('-uploaded_at')
    return render(request, 'subpages/upload_history.html', {'documents': documents})


@login_required
def history_page(request):
    medical_events = MedicalEvent.objects.filter(patient=request.user.patientprofile).order_by('-event_date',
                                                                                               '-created_at')
    return render(request, 'main/history.html', {'medical_events': medical_events})


@login_required
def doctors_page(request):
    return render(request, 'subpages/doctors.html')


@login_required
def profile_page(request):
    return render(request, 'subpages/profile.html')


@login_required
def upload_review_page(request):
    ocr_text = request.session.get('ocr_text', '')
    file_name = request.session.get('uploaded_file_name', '')
    file_url = request.session.get('uploaded_file_url', '')
    selected_options_json = request.session.get('selected_options_json', '{}')

    event_date_missing = request.session.get('event_date_missing', False)

    try:
        selected_options = json.loads(selected_options_json)
    except json.JSONDecodeError:
        selected_options = {}

    context = {
        'ocr_text': ocr_text,
        'file_name': file_name,
        'file_url': file_url,
        'selected_event_type': selected_options.get('event_type_title', _('Не е избран')),
        'selected_category': selected_options.get('category_name', _('Не е избран')),
        'selected_specialty': selected_options.get('specialty_name', _('Не е избран')),
        'selected_doctor': selected_options.get('practitioner_name', _('Не е избран')),
        'selected_event_date': selected_options.get('event_date', _('Не е избрана')),
        'event_date_missing': event_date_missing,
    }
    return render(request, 'subpages/upload_review.html', context)


# --- API Endpoints ---

@csrf_exempt
@login_required
def perform_ocr(request):
    if request.method == 'POST':
        if 'file' not in request.FILES:
            return JsonResponse({'status': 'error', 'message': _('Не е намерен файл за качване.')}, status=400)

        uploaded_file = request.FILES['file']
        file_extension = os.path.splitext(uploaded_file.name)[1].lower()

        if uploaded_file.size > 10 * 1024 * 1024:  # 10 MB limit
            return JsonResponse({'status': 'error', 'message': _('Файлът е твърде голям (макс. 10MB).')}, status=400)

        if file_extension not in ALLOWED_IMAGE_EXTENSIONS and file_extension not in ALLOWED_PDF_EXTENSIONS:
            return JsonResponse({'status': 'error', 'message': _('Неподдържан файлов формат.')}, status=400)

        file_hash = hashlib.sha256(uploaded_file.read()).hexdigest()
        uploaded_file.seek(0)

        existing_document = Document.objects.filter(
            patient=request.user.patientprofile,
            file_hash=file_hash
        ).first()

        if existing_document:
            medical_event_id = None
            try:
                medical_event_id = existing_document.medicalevent.id
            except MedicalEvent.DoesNotExist:
                pass

            return JsonResponse({
                'status': 'success',
                'message': _('Този документ вече е качен и обработен.'),
                'ocr_text': existing_document.ocr_text if existing_document.ocr_text else '',
                'document_id': existing_document.id,
                'medical_event_id': medical_event_id,
                'file_name': uploaded_file.name,
                'file_url': existing_document.file.url,
                'already_processed': True,
            })

        document = Document(
            patient=request.user.patientprofile,
            file=uploaded_file,
            file_hash=file_hash,
        )
        document.save()

        try:
            ocr_text = perform_ocr_space(document.file)
            if not ocr_text:
                document.processing_error_message = _("Неуспешно извличане на текст чрез OCR.")
                document.save()
                return JsonResponse({'status': 'error', 'message': _('OCR обработката не успя да извлече текст.')},
                                    status=500)

            anonymized_ocr_text = anonymizer.anonymize_text(ocr_text)

            request.session['ocr_text'] = anonymized_ocr_text
            request.session['temp_document_id'] = document.id
            request.session['uploaded_file_name'] = uploaded_file.name
            request.session['uploaded_file_url'] = document.file.url

            return JsonResponse({
                'status': 'success',
                'message': _('OCR обработката приключи успешно.'),
                'ocr_text': anonymized_ocr_text,
                'document_id': document.id,
                'file_name': uploaded_file.name,
                'file_url': document.file.url,
            })
        except Exception as e:
            logger.error(f"Error during OCR processing: {e}", exc_info=True)
            document.processing_error_message = f"OCR Error: {str(e)}"
            document.save()
            return JsonResponse({'status': 'error', 'message': _(f'Възникна грешка по време на OCR обработка: {e}')},
                                status=500)
    return JsonResponse({'status': 'error', 'message': _('Невалиден метод на заявка.')}, status=405)


@csrf_exempt
@login_required
def analyze_document(request):
    if request.method == 'POST':
        ocr_text_to_analyze = request.POST.get('edited_ocr_text', request.session.get('ocr_text', ''))
        temp_document_id = request.session.get('temp_document_id')

        if not ocr_text_to_analyze or not temp_document_id:
            return JsonResponse({'status': 'error', 'message': _('Няма текст за анализ или ID на документ.')},
                                status=400)

        event_type_title = request.POST.get('event_type_title')
        category_name = request.POST.get('category_name')
        specialty_name = request.POST.get('specialty_name')
        practitioner_name = request.POST.get('practitioner_name')

        # event_date is now only taken from POST if manually provided on review page (after initial GPT attempt)
        event_date_from_post = request.POST.get('event_date')

        request.session['selected_options_json'] = json.dumps({
            'event_type_title': event_type_title,
            'category_name': category_name,
            'specialty_name': specialty_name,
            'practitioner_name': practitioner_name,
            'event_date': event_date_from_post,  # Store date from POST if available for review page
        })

        try:
            document = get_object_or_404(Document, pk=temp_document_id, patient=request.user.patientprofile)
            document.ocr_text = ocr_text_to_analyze
            document.save()

            gpt_data = call_gpt_for_document(
                ocr_text_to_analyze,
                category_name if category_name else 'general',
                {}
            )

            final_event_type_title = event_type_title if event_type_title else gpt_data.get('event_type_title', 'Other')

            # Prioritize date from POST (manual input) over GPT's extraction
            final_event_date_str = event_date_from_post if event_date_from_post else gpt_data.get('event_date')

            final_event_date = None
            if final_event_date_str:
                try:
                    final_event_date = datetime.strptime(final_event_date_str, '%Y-%m-%d').date()
                except ValueError:
                    final_event_date = None

                    # --- Логика за липсваща дата ---
            if not final_event_date:
                request.session['ocr_text'] = ocr_text_to_analyze
                request.session['temp_document_id'] = document.id
                request.session['uploaded_file_name'] = document.file.name
                request.session['uploaded_file_url'] = document.file.url

                request.session['event_date_missing'] = True

                request.session['selected_options_json'] = json.dumps({
                    'event_type_title': event_type_title,
                    'category_name': category_name,
                    'specialty_name': specialty_name,
                    'practitioner_name': practitioner_name,
                    'event_date': '',
                })

                return JsonResponse({
                    'status': 'redirect_to_review',
                    'message': _('Датата на събитието не беше открита в документа. Моля, въведете я ръчно.'),
                    'redirect_url': reverse('medj:upload_review')
                })
            # --- Край на логиката за липсваща дата ---

            summary = gpt_data.get('summary', '')
            diagnosis = gpt_data.get('diagnosis', '')
            treatment_plan = gpt_data.get('treatment_plan', '')

            final_detected_specialty = specialty_name if specialty_name else gpt_data.get('detected_specialty', '')
            final_category = category_name if category_name else gpt_data.get('category', '')

            suggested_tags = gpt_data.get('suggested_tags', [])

            gpt_practitioner = gpt_data.get('practitioner', {})
            final_practitioner_name = practitioner_name if practitioner_name else gpt_practitioner.get('name')
            final_practitioner_specialty = gpt_practitioner.get('specialty')

            blood_test_results_list = gpt_data.get('blood_test_results', [])
            narrative_sections_list = gpt_data.get('narrative_sections', [])

            with transaction.atomic():
                category_obj = None
                if final_category:
                    category_obj, _ = MedicalCategory.objects.get_or_create(name=final_category)

                specialty_obj = None
                if final_detected_specialty:
                    specialty_obj, _ = MedicalSpecialty.objects.get_or_create(name=final_detected_specialty)

                medical_event_obj, created = MedicalEvent.objects.update_or_create(
                    patient=request.user.patientprofile,
                    source_document=document,
                    defaults={
                        'event_type_title': final_event_type_title,
                        'event_date': final_event_date,
                        'summary': summary,
                        'diagnosis': diagnosis,
                        'treatment_plan': treatment_plan,
                        'specialty': specialty_obj,
                        'category': category_obj,
                    }
                )

                if final_practitioner_name:
                    practitioner, _ = Practitioner.objects.get_or_create(
                        name=final_practitioner_name,
                        defaults={'specialty': specialty_obj}
                    )
                    medical_event_obj.practitioners.add(practitioner)

                medical_event_obj.tags.clear()
                for tag_name in suggested_tags:
                    tag, _ = DocumentTag.objects.get_or_create(name=tag_name)
                    medical_event_obj.tags.add(tag)

                for res_data in blood_test_results_list:
                    indicator_name = res_data.get('indicator_name')
                    value = res_data.get('value')
                    unit = res_data.get('unit', '')
                    reference_range = res_data.get('reference_range', '')

                    try:
                        value = str(value)
                    except (ValueError, TypeError):
                        value = ""

                    if indicator_name and value:
                        lab_indicator_obj, created_indicator = LabIndicator.objects.get_or_create(
                            patient=request.user.patientprofile,
                            indicator_name__iexact=indicator_name,
                            unit__iexact=unit,
                            defaults={
                                'indicator_name': indicator_name,
                                'unit': unit,
                                'reference_range': reference_range
                            }
                        )

                        if not created_indicator and not lab_indicator_obj.reference_range and reference_range:
                            lab_indicator_obj.reference_range = reference_range
                            lab_indicator_obj.save()

                        BloodTestMeasurement.objects.update_or_create(
                            medical_event=medical_event_obj,
                            indicator=lab_indicator_obj,
                            defaults={'value': value}
                        )

                        logger.info(
                            f"Processed BloodTestMeasurement: {indicator_name}: {value} {unit} for MedicalEvent ID {medical_event_obj.id}")

                NarrativeSectionResult.objects.filter(medical_event=medical_event_obj).delete()
                for section_data in narrative_sections_list:
                    section_title = section_data.get('section_title')
                    section_text = section_data.get('section_text')
                    if section_title and section_text:
                        NarrativeSectionResult.objects.create(
                            medical_event=medical_event_obj,
                            section_title=section_title,
                            section_text=section_text
                        )
                        logger.info(
                            f"Created new NarrativeSectionResult: {section_title} for MedicalEvent ID {medical_event_obj.id}")

                if 'ocr_text' in request.session:
                    del request.session['ocr_text']
                if 'temp_document_id' in request.session:
                    del request.session['temp_document_id']
                if 'uploaded_file_name' in request.session:
                    del request.session['uploaded_file_name']
                if 'uploaded_file_url' in request.session:
                    del request.session['uploaded_file_url']
                if 'selected_options_json' in request.session:
                    del request.session['selected_options_json']
                if 'event_date_missing' in request.session:
                    del request.session['event_date_missing']

                return JsonResponse({
                    'status': 'success',
                    'message': _('Документът е анализиран и медицинското събитие е създадено/актуализирано успешно!'),
                    'redirect_url': reverse('medj:event_detail', args=[medical_event_obj.id])
                })

        except json.JSONDecodeError as e:
            logger.error(f"Failed to decode JSON from GPT response: {e}")
            return JsonResponse({'status': 'error', 'message': _(f'Невалиден формат на отговор от GPT: {e}')},
                                status=400)
        except Exception as e:
            logger.error(f"Error during document analysis: {e}", exc_info=True)
            return JsonResponse({'status': 'error', 'message': _(f'Възникна грешка по време на анализ: {e}')},
                                status=500)
    return JsonResponse({'status': 'error', 'message': _('Невалиден метод на заявка.')}, status=405)


@csrf_exempt
@login_required
def get_specialties_for_category(request):
    specialties = MedicalSpecialty.objects.all().order_by('name')
    specialty_list = [{'id': specialty.id, 'name': specialty.name} for specialty in specialties]
    return JsonResponse({'specialties': specialty_list})


@login_required
def event_detail_page(request, event_id):
    medical_event = get_object_or_404(MedicalEvent, pk=event_id, patient=request.user.patientprofile)

    associated_tags = medical_event.tags.all()
    associated_practitioners = medical_event.practitioners.all()

    blood_tests = medical_event.blood_test_measurements.select_related('indicator').order_by(
        'indicator__indicator_name', 'indicator__unit')
    narrative_sections = medical_event.narrative_sections.all()

    context = {
        'medical_event': medical_event,
        'associated_tags': associated_tags,
        'associated_practitioners': associated_practitioners,
        'blood_tests': blood_tests,
        'narrative_sections': narrative_sections,
        'form': MedicalEventForm(instance=medical_event),
        'document': medical_event.source_document
    }
    return render(request, 'subpages/event_detail.html', context)


@csrf_exempt
@login_required
def update_event_details(request, event_id):
    if request.method == 'POST':
        medical_event = get_object_or_404(MedicalEvent, pk=event_id, patient=request.user.patientprofile)

        try:
            data = json.loads(request.body)
            summary = data.get('summary')
            event_date_str = data.get('event_date')
            tags = data.get('tags', [])

            if summary is not None:
                medical_event.summary = summary
            if event_date_str:
                medical_event.event_date = datetime.strptime(event_date_str, '%Y-%m-%d').date()

            medical_event.save()

            medical_event.tags.clear()
            for tag_name in tags:
                tag, _ = DocumentTag.objects.get_or_create(name=tag_name)
                medical_event.tags.add(tag)

            return JsonResponse({'status': 'success', 'message': _('Медицинското събитие е актуализирано успешно.')})
        except json.JSONDecodeError:
            return JsonResponse({'status': 'error', 'message': _('Невалиден JSON формат.')}, status=400)
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)}, status=500)
    return JsonResponse({'status': 'error', 'message': _('Невалиден метод на заявка.')}, status=405)


@csrf_exempt
@login_required
def delete_document(request, document_id):
    if request.method == 'POST':
        document = get_object_or_404(Document, pk=document_id, patient=request.user.patientprofile)

        try:
            medical_event = MedicalEvent.objects.get(source_document=document)
            medical_event.delete()
        except MedicalEvent.DoesNotExist:
            pass

        document.delete()
        return JsonResponse(
            {'status': 'success', 'message': _('Документът и свързаното медицинско събитие са изтрити успешно.')})
    return JsonResponse({'status': 'error', 'message': _('Невалиден метод на заявка.')}, status=405)


@login_required
def export_medical_events_excel(request):
    medical_events = MedicalEvent.objects.filter(patient=request.user.patientprofile).order_by('-event_date')

    wb = Workbook()
    ws = wb.active
    ws.title = _("Медицински Събития")

    headers = [
        _("Дата на Събитието"), _("Тип"), _("Категория"), _("Специалност"),
        _("Обобщение"), _("Диагноза"), _("План за Лечение"), _("Тагове"),
        _("Практикуващи Лекари"), _("Име на Документ"), _("Създадено на"), _("Актуализирано на")
    ]
    ws.append(headers)

    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font

    for event in medical_events:
        event_data = [
            event.event_date.strftime('%Y-%m-%d') if event.event_date else '',
            event.get_event_type_title_display(),
            event.category.name if event.category else '',
            event.specialty.name if event.specialty else '',
            event.summary,
            event.diagnosis,
            event.treatment_plan,
            ', '.join([tag.name for tag in event.tags.all()]),
            ', '.join([p.name for p in event.practitioners.all()]),
            event.source_document.file.name.split('/')[
                -1] if event.source_document and event.source_document.file else '',
            event.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            event.updated_at.strftime('%Y-%m-%d %H:%M:%S'),
        ]
        ws.append(event_data)

        blood_measurements = event.blood_test_measurements.select_related('indicator').order_by(
            'indicator__indicator_name')
        for bm in blood_measurements:
            ws.append([
                '', '', '', '', '',
                _("Кръвен Тест:"),
                bm.indicator.indicator_name,
                f"{bm.value} {bm.indicator.unit}",
                _("Референции:"),
                bm.indicator.reference_range
            ])

        narrative_sections = event.narrative_sections.all().order_by('section_title')
        for ns in narrative_sections:
            ws.append([
                '', '', '', '', '',
                _("Секция:"),
                ns.section_title,
                ns.section_text
            ])

    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if cell.value is not None and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        if adjusted_width > 50:
            adjusted_width = 50
        ws.column_dimensions[column_letter].width = adjusted_width

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="медицински_събития.xlsx"'
    wb.save(response)

    return response


def test_upload_view(request):
    if request.method == 'POST':
        print("\n--- DEBUGGING TEST_UPLOAD_VIEW ---")
        print(f"DEBUG: Test Request method: {request.method}")
        print(f"DEBUG: Test request.FILES content: {request.FILES}")
        print(f"DEBUG: Test request.FILES.get('test_file'): {request.FILES.get('test_file')}")

        if 'test_file' in request.FILES:
            uploaded_file = request.FILES['test_file']
            save_path = os.path.join(settings.MEDIA_ROOT, 'test_uploads', uploaded_file.name)
            os.makedirs(os.path.dirname(save_path), exist_ok=True)
            with open(save_path, 'wb+') as destination:
                for chunk in uploaded_file.chunks():
                    destination.write(chunk)
            print(f"DEBUG: Test file saved to: {save_path}")
            return HttpResponse("File uploaded successfully to test_uploads!", status=200)
        else:
            print("DEBUG: No 'test_file' found in request.FILES")
            return HttpResponse("No file uploaded for test.", status=400)
    return render(request, 'temp_test_upload.html')